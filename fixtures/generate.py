#!/usr/bin/env python3
"""Rebuild synthetic fixtures. Requires Python pillow, python-docx and reportlab.
No patient data, network access or model credentials are used.
"""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parent
LINES = [
    'SYNTHETIC REPORT - TEST DATA ONLY',
    'Name: Synthetic Example',
    'Phone: 202-555-0199',
    'Email: synthetic@example.invalid',
    'Glucose 6.2 mmol/L 3.9-6.1',
    'Hemoglobin 140 g/L 115-150',
    'Platelets 210 10^9/L 125-350',
    'Protein negative',
    'ReviewMarker 7 widgets 1-5',
    'No patient data. Values are fabricated for software verification.',
]
(ROOT/'synthetic-report.txt').write_text('\n'.join(LINES)+'\n', encoding='utf-8')

pdf = canvas.Canvas(str(ROOT/'synthetic-text.pdf'), pagesize=A4)
pdf.setTitle('Synthetic report fixture - no patient data')
for page in [1, 2]:
    pdf.setFont('Helvetica', 13)
    y = 790
    for line in (LINES if page == 1 else ['SYNTHETIC REPORT PAGE 2', 'Sodium 142 mmol/L 137-147', 'Potassium 4.1 mmol/L 3.5-5.3']):
        pdf.drawString(40, y, line)
        y -= 30
    pdf.showPage()
pdf.save()

doc = Document()
doc.add_heading('Synthetic report fixture', level=1)
doc.add_paragraph('TEST DATA ONLY. No patient data. Name: Synthetic Example')
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['Item', 'Result', 'Unit', 'Reference']): cell.text = text
for row in [('Glucose','6.2','mmol/L','3.9-6.1'),('Hemoglobin','140','g/L','115-150'),('Protein','negative','','negative')]:
    for cell, text in zip(table.add_row().cells, row): cell.text = text
doc.add_paragraph('All entries are fabricated for automated software testing.')
doc.save(ROOT/'synthetic-table.docx')

font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
font = ImageFont.truetype(font_path, 38)
img = Image.new('RGB', (1900, 1000), 'white')
draw = ImageDraw.Draw(img)
for i, line in enumerate([LINES[0], LINES[4], LINES[5], 'Protein negative', 'END OF SYNTHETIC FIXTURE']):
    draw.text((80, 100+i*150), line, fill='black', font=font)
img.save(ROOT/'synthetic-scan.png', dpi=(150,150))
img.save(ROOT/'synthetic-scan.pdf', 'PDF', resolution=150, title='Synthetic image-only report fixture')
cjk = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 48)
cjk_img = Image.new('RGB', (1900, 1000), 'white')
cjk_draw = ImageDraw.Draw(cjk_img)
for i, line in enumerate(['合成检验报告  仅供软件测试', '葡萄糖  6.2 mmol/L  3.9-6.1', '血红蛋白  140 g/L  115-150', '此文件不包含真实患者信息']):
    cjk_draw.text((80, 100+i*180), line, fill='black', font=cjk)
cjk_img.save(ROOT/'synthetic-chinese-scan.png', dpi=(150,150))

# Safety regression: tiny ZIP with a deliberately excessive expansion ratio.
# It is never executed and contains only repeated synthetic XML text.
import zipfile
with zipfile.ZipFile(ROOT/'synthetic-expansion-limit.docx', 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('word/document.xml', '<document>' + 'x' * (2 * 1024 * 1024) + '</document>')

print('Generated 7 synthetic report fixtures')
